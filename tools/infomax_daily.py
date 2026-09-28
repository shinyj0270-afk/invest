"""Import saved Infomax daily XLSX values; never executes Excel or fetches quotes."""
import argparse
from datetime import datetime, timezone
import hashlib
import json
from pathlib import Path
import re
import sys

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
sys.path.insert(0, str(ROOT / 'tools'))
from investment.core import validate_snapshot
from tools.infomax_import import METRICS, day, number, require, write_new
from tools.infomax_user_policy import apply_observed_flow_policy, apply_analysis_basis

HEADERS = ['일자', '현재가', '누적거래량', '누적거래대금', '시가총액',
           '기관순매수금액', '외국인순매수금액']


def build(values, formulas, identities, as_of, decisions, provenance):
    as_of = day(as_of, '수집일')
    require(len(values) >= 23, '최소 20개 과거 관측행 필요')
    require(values[0][0] == '시작' and values[0][2] == '종료', '조회 조건행 누락')
    require([str(values[0][i]) for i in (7, 9, 11, 13)] == ['일', 'D', '0', '종가'],
            '일/내림차순/거래일/종가 설정 필요')
    require(day(values[0][3], '종료일') == as_of, '수집일과 조회 종료일 불일치')
    header = list(values[2])
    while header and header[-1] is None:
        header.pop()
    require(header and len(header) % len(HEADERS) == 0, '7열 종목 블록 필요')
    companies, histories, excluded = [], {}, {}
    for col in range(0, len(header), len(HEADERS)):
        require(header[col:col+7] == HEADERS, '필드 순서/단위 계약 불일치')
        formula = formulas[1][col]
        match = re.match(r'^=(?:_xll\.)?IMDH\("STK","(\d{6})",', str(formula))
        require(match is not None, '종목코드는 원본 IMDH 함수에서 확인해야 합니다')
        code = match[1]
        require(code in identities and code not in histories, '미등록 또는 중복 종목코드')
        identity = identities[code]
        require(values[1][col] == identity['name'], '종목 이름과 코드 불일치')
        observations, dates, omitted = {}, [], []
        for raw in values[3:]:
            cells = list(raw[col:col+7])
            if not any(v is not None for v in cells):
                continue
            require(len(cells) == 7, '잘린 종목 블록')
            observed = day(cells[0], '관측일')
            require(observed <= as_of and observed not in dates, '미래 또는 중복 관측일')
            dates.append(observed)
            if observed == as_of:
                omitted.append(observed)
                continue
            vals = [number(v, HEADERS[i+1]) for i, v in enumerate(cells[1:])]
            require(all(v is not None for v in vals), '과거 관측 결측: 기간을 줄여 대체하지 않습니다')
            require(vals[0] > 0 and vals[1] >= 0 and vals[2] >= 0 and vals[3] > 0,
                    '가격/거래량/거래대금/시가총액 범위 오류')
            observations[observed] = dict(zip(HEADERS[1:], vals))
        require(dates == sorted(dates, reverse=True), '내림차순 날짜 정렬 필요')
        require(len(observations) >= 20, '당일 제외 후 과거 관측 20개 미만')
        histories[code] = observations
        excluded[code] = omitted
        companies.append(dict(code=code, **identity))
    require(set(histories) == set(identities), '설정과 실제 종목 집합 불일치')
    sessions = sorted(next(iter(histories.values())))
    require(all(sorted(h) == sessions for h in histories.values()), '종목별 관측일 불일치')
    end, window = sessions[-1], sessions[-20:]
    warnings = ['거래소 범위 미확인', '추후 정정 가능', '공식 거래일 미대조',
                '공급자 분류 기준·세부 범위 미확인', '가격·수급 거래소 범위 동일성 미확인',
                '재무자료 미수집', '일별 현재가는 수정주가로 검증되지 않음',
                '업종은 종목 설정의 수동 분류이며 공급자 업종/GICS가 아님']
    for row in companies:
        code = row['code']; hist = histories[code]
        row.update(metrics={k: None for k in METRICS}, history=[], annual=[], quarters=[], evidence=[],
                   stages=[], price_venue=None, data_quality=warnings.copy(),
                   sources=[dict(label='이 PC 인포맥스 1333 일별 XLSX 저장본', url='',
                                 file=provenance['file'], sha256=provenance['sha256'])],
                   metric_missing_reasons={k: '필요한 재무 원자료·공개 시점 미수집' for k in METRICS})
        for key in METRICS:
            if '20d' in key:
                row['metric_missing_reasons'][key] = '가격·수급의 잠정 사용 정책 또는 원자료 검증 필요'
        row['metrics'].update(price=hist[end]['현재가'])
        # Preserve the existing explicit policy for provisional market observations.
        if (decisions.get('unknown_venue', {}).get('decision') == 'allow_use_with_label' and
            decisions.get('revision_finality', {}).get('decision') == 'allow_provisional_use_with_label'):
            row['metrics']['market_cap_eok'] = hist[end]['시가총액'] / 1e8
            row['metric_missing_reasons'].pop('market_cap_eok', None)
        else:
            row['metric_missing_reasons']['market_cap_eok'] = '거래소·확정 여부 검증 필요'
        apply_observed_flow_policy(row, dict(prices=histories, flows=histories), window, decisions)
        row['prices'] = [dict(date=d, close=hist[d]['현재가'], volume=hist[d]['누적거래량'],
                              turnover=hist[d]['누적거래대금'], venue=None, final=False,
                              adjustment_basis='unverified') for d in sessions]
        row['flows'] = [dict(date=d, institution_net_won=hist[d]['기관순매수금액']*1000,
                             foreign_net_won=hist[d]['외국인순매수금액']*1000,
                             venue=None, final=False) for d in sessions]
    snapshot = dict(schema_version='0.1', meta=dict(
        data_mode='user_input', as_of=as_of, price_date=end, flow_start=window[0], flow_end=end,
        financial_period='미수집', financial_basis='CFS',
        financial_basis_note='향후 재무 입력 계약이며 이번 파일에는 재무자료가 없음',
        venue='가격·수급 거래소 미확인', universe_label='집 PC 인포맥스 일별 조회 표본',
        universe_total=None, calendar_basis='observed_dates_unverified',
        adapter_version='infomax-daily-0.1', source_files=provenance, warnings=warnings,
        fetched_at=datetime.now(timezone.utc).isoformat(),
        import_audit=dict(observations_per_company=len(sessions), excluded_same_day=excluded,
                          units={'turnover':'원','market_cap':'원','flow':'천원→원 ×1000'})),
        companies=companies, sessions=sessions, benchmarks={})
    return validate_snapshot(apply_analysis_basis(snapshot, decisions))


def convert(path, config, as_of=None):
    from openpyxl import load_workbook
    path = Path(path)
    require(path.stat().st_size < 30*1024*1024, '입력 파일 크기 제한 초과')
    original = path.read_bytes()
    from io import BytesIO
    cached = load_workbook(BytesIO(original), read_only=True, data_only=True)
    formulas = load_workbook(BytesIO(original), read_only=True, data_only=False)
    try:
        require(len(cached.sheetnames) == 1, '한 시트 입력 필요')
        require(cached.active.max_row <= 10005 and cached.active.max_column <= 700, '입력 크기 제한 초과')
        values = list(cached.active.values)
        formula_values = [[getattr(v, 'text', v) for v in row] for row in formulas.active.values]
        return build(values, formula_values, config['companies'],
                     as_of or day(values[0][3], '저장된 조회 종료일'), config['decisions'],
                     dict(file=path.name, sha256=hashlib.sha256(original).hexdigest()))
    finally:
        cached.close(); formulas.close()


def main():
    p = argparse.ArgumentParser(description=__doc__)
    p.add_argument('--input', required=True, type=Path)
    p.add_argument('--config', required=True, type=Path)
    p.add_argument('--as-of', required=True)
    p.add_argument('--output', required=True, type=Path)
    args = p.parse_args()
    try:
        config = json.loads(args.config.read_text(encoding='utf-8-sig'))
        snapshot = convert(args.input, config, args.as_of)
        write_new(args.output, snapshot)
    except (ValueError, OSError, KeyError, TypeError) as exc:
        p.exit(2, f'변환 실패: {exc}\n')
    print(json.dumps(dict(companies=len(snapshot['companies']), price_date=snapshot['meta']['price_date'],
                          flow_start=snapshot['meta']['flow_start'], mode='user_input'), ensure_ascii=False))


if __name__ == '__main__':
    main()
