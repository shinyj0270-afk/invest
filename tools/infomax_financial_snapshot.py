"""Join saved Infomax quarterly accounts to a daily snapshot; no Excel/API execution.

The review JSON is an explicit, dated account-scope review, not an inferred
provider-wide accounting definition. All monetary input cells are KRW thousands.
"""
import argparse
import copy
from datetime import datetime, timezone
import hashlib
from io import BytesIO
import json
from pathlib import Path
import re
import sys
from zipfile import ZipFile

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))
from investment.core import continuous, ratio, validate_snapshot
from tools.infomax_import import day, number, require, write_new

BALANCE = ['일자', '자산', '자본', '부채', '현금및현금성자산', '기말비지배주주지분',
           '유동자산(요약재무)', '유동부채(요약재무)', '차입금', '사채',
           '단기차입금(요약재무)', '장기차입금(요약재무)', '유동성장기부채(요약재무)',
           '리스부채(요약재무)', '유동성리스부채(요약재무)', '단기사채(요약재무)',
           '사채(요약재무)', '지배기업주주지분(요약재무)', '비지배주주지분(요약재무)']
INCOME = ['일자', '매출액(영업수익)', '영업이익', '당기순이익지배기업주주지분',
          '이자비용', '당기순이익(포괄손익계산서)']
FINANCIAL_METRICS = ['revenue_growth_pct', 'operating_margin_pct', 'roe_pct',
                     'debt_ratio_pct', 'net_debt_equity_pct', 'interest_coverage_x',
                     'current_ratio_pct']
CORE_FIELDS = {'assets': '자산', 'equity': '자본', 'liabilities': '부채',
               'cash': '현금및현금성자산', 'nci': '기말비지배주주지분',
               'current_assets': '유동자산(요약재무)', 'current_liabilities': '유동부채(요약재무)',
               'revenue': '매출액(영업수익)', 'profit': '영업이익',
               'parent_income': '당기순이익지배기업주주지분', 'interest': '이자비용'}


def parse(values, formulas, fields, identities, as_of, comp):
    require(len(values) >= 8, '최소 5분기 필요')
    require(values[0][0] == '시작' and values[0][2] == '종료', '조회 조건행 필요')
    require(day(values[0][3], '종료') == as_of, '재무 조회 종료일 불일치')
    require(values[0][7] == '분기' and values[0][9] == 'D', '분기·내림차순 필요')
    width = len(fields)
    header = list(values[2])
    while header and header[-1] is None:
        header.pop()
    require(len(header) == width * len(identities), '종목 블록 수 불일치')
    result = {}
    for col in range(0, len(header), width):
        require(header[col:col+width] == fields, '계정 이름·순서 불일치')
        formula = str(formulas[1][col])
        match = re.match(r'^=(?:_xll\.)?IMDH\("STK","(\d{6})",', formula)
        require(match is not None, '원본 IMDH 종목코드 필요')
        code = match[1]
        require(code in identities and code not in result, '미등록·중복 종목')
        require(values[1][col] == identities[code], '종목명 불일치')
        # Exact option delimiters prevent 연결우선 / cumulative income substitutions.
        for option in ['Per=분기', 'sort=D', 'real=false', 'Orient=V', 'Cons=연결',
                       f'Comp={comp}', 'Trai=연속', 'unit=true']:
            require(re.search(r'(?:[",])' + re.escape(option) + r'(?:[",])', formula),
                    f'재무 함수 조건 불일치: {option}')
        rows = {}
        for raw in values[3:]:
            cells = list(raw[col:col+width])
            if not any(v is not None for v in cells):
                continue
            require(len(cells) == width, '잘린 계정 행')
            end = day(cells[0], '분기말')
            require(end <= as_of and end not in rows, '미래·중복 분기')
            rows[end] = {f: number(v, f) for f, v in zip(fields[1:], cells[1:])}
        dates = list(rows)
        require(dates == sorted(dates, reverse=True), '분기 정렬 오류')
        require(len(dates) >= 5 and continuous([{'period_end': d} for d in reversed(dates)], len(dates)),
                '비연속 분기 또는 전년 비교분기 누락')
        result[code] = rows
    require(set(result) == set(identities), '종목 집합 불일치')
    return result


def read_accounts(path, fields, identities, as_of, comp):
    from openpyxl import load_workbook
    path = Path(path)
    require(path.stat().st_size <= 30*1024*1024, '입력 크기 초과')
    payload = path.read_bytes()
    with ZipFile(BytesIO(payload)) as archive:
        require(sum(x.file_size for x in archive.infolist()) <= 100*1024*1024, '압축 해제 크기 초과')
    cached = load_workbook(BytesIO(payload), read_only=True, data_only=True, keep_links=False)
    formulas = load_workbook(BytesIO(payload), read_only=True, data_only=False, keep_links=False)
    try:
        require(len(cached.sheetnames) == 1, '단일 시트 필요')
        ws = cached.active
        require(ws.max_row <= 1000 and ws.max_column <= 512, '시트 크기 초과')
        fs = [[getattr(v, 'text', v) for v in row] for row in formulas.active.values]
        data = parse(list(ws.values), fs, fields, identities, as_of, comp)
    finally:
        cached.close()
        formulas.close()
    return data, dict(file=path.name, sha256=hashlib.sha256(payload).hexdigest(), unit='천원')


def calculate(balance, income, review):
    dates = list(income)
    require(set(balance) == set(income), '재무상태표·손익 분기 불일치')
    latest, previous = dates[0], dates[4]
    require(review['period_end'] == latest, '계정 범위 검토 기간 불일치')
    b, opening, inc = balance[latest], balance[previous], income[latest]
    for end, row in balance.items():
        for f in ['자산', '자본', '부채']:
            require(row[f] is not None and row[f] >= 0, f'{end}: 필수 잔액 결측·음수')
        require(abs(row['자산']-row['자본']-row['부채'])*1000 < 100_000_000,
                f'{end}: 자산=자본+부채 불일치')
    for row in (b, opening):
        require(row['기말비지배주주지분'] is not None, '비지배지분 결측')
        require(row['지배기업주주지분(요약재무)'] is not None, '지배기업 지분 결측')
        require(abs(row['자본']-row['기말비지배주주지분']-row['지배기업주주지분(요약재무)'])*1000 < 100_000_000,
                '지배기업 지분 분모 대조 실패')
    for d in dates[:5]:
        require(all(income[d][f] is not None for f in INCOME[1:]), '손익 필수 계정 결측')
    for f in ['현금및현금성자산', '유동자산(요약재무)', '유동부채(요약재무)']:
        require(b[f] is not None and b[f] >= 0, f'{f}: 필수 잔액 결측·음수')
    parent_ttm = sum(income[d]['당기순이익지배기업주주지분'] for d in dates[:4])
    avg_equity = ((b['자본']-b['기말비지배주주지분']) +
                  (opening['자본']-opening['기말비지배주주지분'])) / 2
    # A reviewed partition is mandatory. A blank lease field is never treated as zero.
    components = review['debt_fields']
    require(len(components) == len(set(components)) and components, '차입금 구성 중복·누락')
    require(all(f in BALANCE[1:] and b[f] is not None and b[f] >= 0 for f in components),
            '차입금 구성 계정 결측·음수')
    require(review.get('debt_scope_note') and review.get('sources'), '차입금 범위 근거 필요')
    require(all(b[f] is not None for f in ['차입금', '단기차입금(요약재무)', '장기차입금(요약재무)']),
            '차입금 분할 대조 계정 결측')
    require(abs(b['차입금']-b['단기차입금(요약재무)']-b['장기차입금(요약재무)'])*1000 < 100_000_000,
            '차입금=단기+장기 대조 실패: 기존 구성 재사용 불가')
    for f, explanation in review.get('omitted_blank_fields', {}).items():
        require(f in BALANCE and b[f] is None and explanation, '빈 항목 제외 근거 불일치')
    debt = sum(b[f] for f in components)
    metrics = dict(revenue_growth_pct=ratio(inc['매출액(영업수익)']-income[previous]['매출액(영업수익)'], income[previous]['매출액(영업수익)'], 100),
        operating_margin_pct=ratio(inc['영업이익'], inc['매출액(영업수익)'], 100),
        roe_pct=ratio(parent_ttm, avg_equity, 100),
        debt_ratio_pct=ratio(b['부채'], b['자본'], 100),
        net_debt_equity_pct=ratio(debt-b['현금및현금성자산'], b['자본'], 100),
        interest_coverage_x=ratio(inc['영업이익'], inc['이자비용']),
        current_ratio_pct=ratio(b['유동자산(요약재무)'], b['유동부채(요약재무)'], 100))
    require(all(v is not None for v in metrics.values()), '재무 비율 분모가 0 이하: 자동 대체 안 함')
    observed = {k: (b if f in b else inc)[f]/1000 for k, f in CORE_FIELDS.items()}
    observed.update(previous_revenue=income[previous]['매출액(영업수익)']/1000,
        opening_equity=opening['자본']/1000, opening_nci=opening['기말비지배주주지분']/1000,
        parent_ttm=parent_ttm/1000, total_debt=debt/1000)
    official = review['official_million_krw']
    require(set(official) == set(observed), '공식 대조 계정 집합 불일치')
    diffs = {k: observed[k]-official[k] for k in observed}
    audit = dict(period_end=latest, previous_comparison=previous, basis='CFS',
        income_basis='standalone_quarter', roe_window=dates[:4],
        observed_million_krw=observed, official_million_krw=official,
        differences_million_krw=diffs,
        differences_at_least_100m=[k for k,v in diffs.items() if abs(v)*1e6 >= 100_000_000],
        debt_components_thousand_krw={f:b[f] for f in components},
        debt_scope_note=review['debt_scope_note'], omitted_blank_fields=review.get('omitted_blank_fields',{}),
        selected_source='infomax', review_sources=review['sources'],
        historical_publication_dates_verified=False)
    return metrics, audit


def build(base, balances, incomes, review, provenance):
    base = copy.deepcopy(validate_snapshot(base))
    require(base['meta']['data_mode'] == 'user_input', '실제 저장자료 스냅샷만 지원')
    codes = {r['code'] for r in base['companies']}
    require(set(balances) == set(incomes) == set(review['companies']) == codes, '재무 종목 집합 불일치')
    require(review['retrieved_on'] == base['meta']['as_of'], '수집 기준일 불일치')
    require(review['source_priority'] == 'infomax', '인포맥스 우선 정책 명시 필요')
    latest = {next(iter(v)) for v in incomes.values()}
    require(len(latest) == 1, '혼합 재무 보고기간')
    notes = ['재무는 연결 기준. 매출증가율·영업이익률·이자보상배율은 단독분기, ROE는 최근 4분기 지배순이익/기초·기말 평균 지배자본.',
             '현재 조회된 재무자료 기준이며 최초 공개일·정정 이력이 검증된 과거 시점 자료가 아님.',
             '순차입금은 검토된 차입금·사채·유동성 부채·리스를 합산하고 현금및현금성자산만 차감. 단기금융상품은 미차감.']
    for row in base['companies']:
        code = row['code']
        metrics, audit = calculate(balances[code], incomes[code], review['companies'][code])
        row['metrics'].update(metrics)
        for key in metrics:
            row['metric_missing_reasons'].pop(key, None)
        row['financial_evidence'] = audit
        audit['notes'] = review['companies'][code].get('notes', [])
        row['financial_accounts'] = dict(balance=balances[code], income=incomes[code], unit='천원')
        # Do not invent release dates to enable historical screens or backtests.
        row['quarters'] = [dict(period_end=d, available_at=None, retrieved_on=review['retrieved_on'],
            basis='CFS', revenue=v['매출액(영업수익)']*1000, op=v['영업이익']*1000,
            parent_income=v['당기순이익지배기업주주지분']*1000)
            for d,v in sorted(incomes[code].items())]
        row['data_quality'] = [n for n in row['data_quality'] if n != '재무자료 미수집']
        row['data_quality'].extend(notes + [audit['debt_scope_note']] + review['companies'][code].get('notes', []))
        row['sources'].extend([dict(label='이 PC 인포맥스 연결 재무 저장본', url='', **source) for source in provenance.values()])
        row['sources'].extend(review['companies'][code]['sources'])
    meta = base['meta']
    meta.update(financial_period=next(iter(latest)), financial_basis_note='인포맥스 연결 잔액·단독분기 손익',
                financial_observed_on=review['retrieved_on'], adapter_version='infomax-daily-financial-0.1')
    meta['financial_source_files'] = provenance
    meta['financial_review'] = dict(source_priority='infomax', amount_tolerance_won=100_000_000,
                                     tolerance_comparison='abs(difference)<threshold', historical_vintages_verified=False)
    meta['warnings'] = [n for n in meta['warnings'] if n != '재무자료 미수집']
    meta['warnings'].append('재무 최초 공개일·정정 이력 미검증: 과거 시점 재무·장기성장 신호는 대기')
    meta['warnings'].extend(review.get('warnings', []))
    meta['fetched_at'] = datetime.now(timezone.utc).isoformat()
    return validate_snapshot(base)


def main():
    p = argparse.ArgumentParser(description=__doc__)
    for arg in ['base', 'balance', 'income', 'review', 'output']:
        p.add_argument('--'+arg, type=Path, required=True)
    args = p.parse_args()
    try:
        base = json.loads(args.base.read_text(encoding='utf-8-sig'))
        review = json.loads(args.review.read_text(encoding='utf-8-sig'))
        identities = {r['code']:r['name'] for r in base['companies']}
        as_of = day(base['meta']['as_of'], '수집일')
        balances, bsource = read_accounts(args.balance, BALANCE, identities, as_of, '누적')
        incomes, isource = read_accounts(args.income, INCOME, identities, as_of, '순')
        result = build(base, balances, incomes, review, dict(balance=bsource, income=isource))
        result['meta']['financial_review']['sha256'] = hashlib.sha256(args.review.read_bytes()).hexdigest()
        write_new(args.output, result)
    except (ValueError, OSError, KeyError, TypeError) as exc:
        p.exit(2, f'재무 연결 실패: {exc}\n')
    print(json.dumps(dict(companies=len(result['companies']), financial_period=result['meta']['financial_period'],
                         financial_metrics_per_company=7), ensure_ascii=False))


if __name__ == '__main__':
    main()
