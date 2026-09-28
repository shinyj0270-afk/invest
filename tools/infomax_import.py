"""Offline reader for the verified Infomax 1333 export layout. No Excel execution."""
import argparse
import csv
import hashlib
import io
import json
import math
import os
import re
import tempfile
from datetime import date, datetime
from pathlib import Path
from zipfile import BadZipFile, ZipFile

VERSION = 'infomax-file-0.1'
FIELDS = {
    'prices': ['현재가', '누적거래량', '누적거래대금'],
    'flows': ['기관순매수금액', '외국인순매수금액'],
    'balance': ['자산', '자본', '부채'],
    'income': ['매출액(영업수익)', '영업이익', '당기순이익(포괄손익계산서)'],
}
METRICS = ['market_cap_eok', 'revenue_growth_pct', 'operating_margin_pct',
           'roe_pct', 'debt_ratio_pct', 'net_debt_equity_pct', 'interest_coverage_x',
           'current_ratio_pct', 'foreign_net_20d_eok', 'institution_net_20d_eok',
           'foreign_net_turnover_20d_pct', 'institution_net_turnover_20d_pct',
           'avg_trading_value_20d_eok']
UNITS = {'누적거래대금': '원', '시가총액': '원',
         **{f: '천원' for k in ('flows', 'balance', 'income') for f in FIELDS[k]}}
LIMIT = 30 * 1024 * 1024


class InputError(ValueError):
    pass


def require(ok, message):
    if not ok:
        raise InputError(message)


def text(value):
    return '' if value is None else str(value).strip()


def day(value, label):
    if isinstance(value, datetime):
        require(value.time() == datetime.min.time(), f'{label}: 시각이 섞인 날짜')
        value = value.date()
    if isinstance(value, date):
        return value.isoformat()
    value = text(value)
    require(bool(re.fullmatch(r'\d{4}-\d{2}-\d{2}', value)), f'{label}: YYYY-MM-DD 필요')
    try:
        return date.fromisoformat(value).isoformat()
    except ValueError as exc:
        raise InputError(f'{label}: 존재하지 않는 날짜') from exc


def number(value, label):
    if value is None or value == '':
        return None
    require(not isinstance(value, bool), f'{label}: boolean은 숫자가 아님')
    if isinstance(value, str):
        value = value.strip()
        if not value:
            return None
        # Refuse Excel errors, formula text, ambiguous dash placeholders and bad grouping.
        require(bool(re.fullmatch(r'[+-]?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.\d+)?(?:[eE][+-]?\d+)?', value)),
                f'{label}: 숫자 또는 빈 셀 필요')
        value = value.replace(',', '')
    try:
        result = float(value)
    except (ValueError, TypeError) as exc:
        raise InputError(f'{label}: 잘못된 숫자') from exc
    require(math.isfinite(result), f'{label}: 유한한 숫자 필요')
    return result


def read_grid(path, sheet=None):
    path = Path(path)
    require(path.stat().st_size <= LIMIT, f'{path.name}: 30MB 제한')
    payload = path.read_bytes()
    if path.suffix.lower() == '.csv':
        try:
            source = payload.decode('utf-8-sig')
        except UnicodeDecodeError:
            source = payload.decode('cp949')
        grid = list(csv.reader(io.StringIO(source)))
    elif path.suffix.lower() == '.xlsx':
        try:
            import openpyxl
        except ImportError as exc:
            raise InputError('XLSX 읽기에는 requirements-import.txt 설치가 필요합니다.') from exc
        with ZipFile(io.BytesIO(payload)) as archive:
            require(sum(i.file_size for i in archive.infolist()) <= 100 * 1024 * 1024,
                    'XLSX 압축 해제 크기 제한 초과')
        wb = openpyxl.load_workbook(io.BytesIO(payload), read_only=True, data_only=True, keep_links=False)
        try:
            require(sheet is not None or len(wb.sheetnames) == 1, '여러 시트이면 --sheet를 지정하세요.')
            require(sheet is None or sheet in wb.sheetnames, '지정한 시트가 없습니다.')
            ws = wb[sheet] if sheet else wb.worksheets[0]
            require((ws.max_row or 0) <= 10000 and (ws.max_column or 0) <= 512, '시트 크기 제한 초과')
            grid = list(ws.iter_rows(values_only=True))
        finally:
            wb.close()
    else:
        raise InputError('CSV 또는 XLSX만 지원합니다. XLS는 Excel에서 XLSX로 저장하세요.')
    require(len(grid) <= 10000 and all(len(r) <= 512 for r in grid), '표 크기 제한 초과')
    return grid, {'file': path.name, 'sha256': hashlib.sha256(payload).hexdigest()}


def at(grid, row, col):
    return grid[row][col] if row < len(grid) and col < len(grid[row]) else None


def parse_info(grid):
    require(bool(grid), '기본정보가 비었습니다.')
    headers = [text(v) for v in grid[0]]
    needed = ['종목명', '구분', '코드', '단축코드', '소속시장구분', '업종코드', '시가총액']
    require(headers[:7] == needed and not any(headers[7:]), '기본정보 헤더가 검증된 1333 형식과 다릅니다.')
    result, names = {}, set()
    for r, row in enumerate(grid[1:], 2):
        if not any(text(v) for v in row):
            continue
        require(len(row) >= 7, f'기본정보 {r}행: 열 누락')
        name, subject, code, short, market, industry, cap = row[:7]
        code = text(code)
        require(bool(re.fullmatch(r'\d{6}', code)), f'기본정보 {r}행: 코드 선행 0을 보존하세요.')
        require(text(subject) == 'STK' and text(short) == code, f'{code}: STK/단축코드 불일치')
        name, industry = text(name), text(industry)
        require(name and industry, f'{code}: 이름/업종 누락 (Excel 계산 후 값으로 저장 필요)')
        require(code not in result and name not in names, '종목코드 또는 이름 중복: 자동 블록 연결 불가')
        require(text(market) in ('1', '7'), f'{code}: 미지원 소속시장 코드')
        cap = number(cap, f'{code} 시가총액')
        require(cap is None or cap > 0, f'{code}: 시가총액은 양수여야 합니다.')
        result[code] = dict(code=code, name=name, market={'1': 'KOSPI', '7': 'KOSDAQ'}[text(market)],
                            industry=industry, cap=cap)
        names.add(name)
    require(0 < len(result) <= 100, '샘플 변환은 1~100종목만 지원합니다.')
    return result


def parse_history(grid, kind, info):
    require(len(grid) >= 3, f'{kind}: 1333 조건/종목명/항목 3행 필요')
    require(text(at(grid, 0, 0)) == '시작', f'{kind}: 조건행 누락')
    require(text(at(grid, 0, 7)) == ('일' if kind in ('prices', 'flows') else '분기'), f'{kind}: 주기 불일치')
    require(text(at(grid, 0, 9)) == 'D' and text(at(grid, 0, 11)) == '0', f'{kind}: 내림차순·영업일 0 필요')
    require(text(at(grid, 0, 13)) == '종가', f'{kind}: 종가 설정 필요')
    width = len(FIELDS[kind]) + 1
    by_name = {v['name']: code for code, v in info.items()}
    result, seen = {}, set()
    for c in range(0, len(grid[2]), width):
        headers = [text(at(grid, 2, c + n)) for n in range(width)]
        if not any(headers):
            continue
        require(headers == ['일자'] + FIELDS[kind], f'{kind}: {c+1}열 헤더 불일치')
        name = text(at(grid, 1, c))
        require(name in by_name and name not in seen, f'{kind}: 종목명 누락/중복/미등록 (값으로 저장 확인)')
        seen.add(name)
        code, observations, order = by_name[name], {}, []
        for r in range(3, len(grid)):
            cells = [at(grid, r, c + n) for n in range(width)]
            if not any(text(v) for v in cells):
                continue
            d = day(cells[0], f'{kind}/{code}/{r+1}행')
            require(d not in observations, f'{kind}/{code}: 날짜 중복 {d}')
            values = {f: number(v, f'{kind}/{code}/{d}/{f}') for f, v in zip(FIELDS[kind], cells[1:])}
            if kind == 'prices':
                require(values['현재가'] is None or values['현재가'] > 0, f'{code}/{d}: 가격은 양수여야 합니다.')
                require(all(values[f] is None or values[f] >= 0 for f in FIELDS[kind][1:]), f'{code}/{d}: 거래량/대금 음수')
            observations[d] = values
            order.append(d)
        require(order == sorted(order, reverse=True), f'{kind}/{code}: 날짜 정렬 오류')
        result[code] = observations
    require(set(result) == set(info), f'{kind}: 기본정보와 종목 집합 불일치')
    return result


def validate_config(config):
    require(isinstance(config, dict), '설정은 JSON 객체여야 합니다.')
    require(config.get('schema_version') == 'infomax-import-0.1', '설정 schema_version 오류')
    require(config.get('mode') in ('fixture', 'user_input'), 'mode는 fixture/user_input')
    as_of = day(config.get('as_of'), 'as_of')
    end = day(config.get('completed_through'), 'completed_through')
    require(end <= as_of, '완료 거래일이 평가일 이후입니다.')
    sessions = config.get('sessions_20d')
    require(isinstance(sessions, list) and len(sessions) == 20, '공식 거래일 20개를 sessions_20d에 지정하세요.')
    sessions = [day(d, '거래일') for d in sessions]
    require(sessions == sorted(set(sessions)) and sessions[-1] == end, '20거래일은 중복 없이 오름차순이며 완료일까지여야 합니다.')
    period = day(config.get('financial_period'), 'financial_period')
    require(period[5:] in ('03-31', '06-30', '09-30', '12-31') and period <= as_of, '유효한 과거 분기 말일 필요')
    require(config.get('financial_basis') == 'CFS', '이번 어댑터는 연결 CFS만 지원합니다.')
    require(config.get('balance_comp') == '누적' and config.get('income_comp') == '순', '잔액=누적, 손익=순 설정 필요')
    require(config.get('units') == UNITS, '단위 설정이 검증된 필드 정의와 다릅니다.')
    for k in ('price_venue', 'flow_venue'):
        require(config.get(k) is None or config[k] in ('KRX', 'NXT', 'KRX+NXT'), f'{k}: 미지원 범위')
    for k in ('prices_final', 'flows_final'):
        require(isinstance(config.get(k), bool), f'{k}: true/false 필요')
    if config.get('market_cap_date') is not None:
        require(day(config['market_cap_date'], 'market_cap_date') <= as_of, '시가총액 날짜가 미래입니다.')
    require(isinstance(config.get('companies'), dict), 'companies 분류/공개일 설정 필요')
    return as_of, end, sessions, period


def build_snapshot(grids, config, provenance=None):
    as_of, end, sessions, period = validate_config(config)
    info = parse_info(grids['info'])
    require(set(config['companies']) == set(info), '설정과 기본정보의 종목 집합 불일치')
    data = {k: parse_history(grids[k], k, info) for k in FIELDS}
    warnings = ['파일 입력이며 공급자 자동 연결·공시 원문 검증은 수행하지 않았습니다.',
                '거래일·확정 여부·연결/누적/순 기준은 사용자 설정이며 파일만으로 독립 검증되지 않습니다.']
    companies, audit = [], {}
    for code, identity in info.items():
        settings = config['companies'][code]
        require(isinstance(settings, dict), f'{code}: 분류 설정 오류')
        security = settings.get('security_type', 'unknown')
        profile = settings.get('analysis_profile', 'unknown')
        require(security in ('ordinary', 'preferred', 'unknown') and profile in ('nonfinancial', 'financial', 'unknown'), f'{code}: 분류 값 오류')
        reasons = {}
        metrics = {k: None for k in METRICS}

        def missing(keys, reason):
            for key in keys:
                reasons[key] = reason

        for kind in data:
            require(all(d <= as_of for d in data[kind][code]), f'{code}/{kind}: 평가일 이후 자료가 있습니다.')
        if config.get('market_cap_date') == end and config['prices_final']:
            metrics['market_cap_eok'] = identity['cap'] / 1e8 if identity['cap'] is not None else None
        if metrics['market_cap_eok'] is None:
            missing(['market_cap_eok'], '시가총액 값·완료 기준일 일치·확정 여부 확인 필요')

        prices = data['prices'][code]
        turnovers = [prices.get(d, {}).get('누적거래대금') for d in sessions]
        valid_turnover = config['prices_final'] and config.get('price_venue') and all(v is not None for v in turnovers)
        total_turnover = sum(turnovers) if valid_turnover else None
        if valid_turnover:
            metrics['avg_trading_value_20d_eok'] = total_turnover / 20 / 1e8
        else:
            missing(['avg_trading_value_20d_eok'], '20거래일 거래대금 누락 또는 가격 확정/거래소 범위 미확인')
        for field, prefix in [('외국인순매수금액', 'foreign'), ('기관순매수금액', 'institution')]:
            amount_key, ratio_key = f'{prefix}_net_20d_eok', f'{prefix}_net_turnover_20d_pct'
            values = [data['flows'][code].get(d, {}).get(field) for d in sessions]
            if config['flows_final'] and config.get('flow_venue') and all(v is not None for v in values):
                total = sum(values) * 1000
                metrics[amount_key] = total / 1e8
                if total_turnover is not None and total_turnover > 0 and config['price_venue'] == config['flow_venue']:
                    metrics[ratio_key] = total / total_turnover * 100
                else:
                    missing([ratio_key], '동일 거래소·동일 20거래일의 양수 거래대금 필요')
            else:
                missing([amount_key, ratio_key], '20거래일 수급 누락 또는 확정/거래소 범위 미확인')

        available = settings.get('financial_available_on')
        if available is not None:
            available = day(available, f'{code} 재무 공개일')
            require(period <= available <= as_of, f'{code}: 재무 공개일이 보고기간/평가일과 모순됩니다.')
        b = data['balance'][code].get(period, {})
        inc = data['income'][code].get(period, {})
        if available:
            revenue, profit = inc.get('매출액(영업수익)'), inc.get('영업이익')
            if revenue is not None and revenue > 0 and profit is not None:
                metrics['operating_margin_pct'] = profit / revenue * 100
            equity, liabilities, assets = b.get('자본'), b.get('부채'), b.get('자산')
            if None not in (assets, equity, liabilities):
                require(abs(assets-equity-liabilities) <= max(1, abs(assets)*1e-8), f'{code}: 자산=자본+부채 불일치')
            if equity is not None and equity > 0 and liabilities is not None and liabilities >= 0:
                metrics['debt_ratio_pct'] = liabilities / equity * 100
        for key in ('operating_margin_pct', 'debt_ratio_pct'):
            if metrics[key] is None:
                missing([key], '대상 분기 원항목·양수 분모·공개일 확인 필요')
        missing(['revenue_growth_pct'], '전년 동분기 매출과 공개 시점 미검증')
        missing(['roe_pct'], 'TTM 지배주주 순이익·기초/기말 지배주주 지분 미수집')
        missing(['net_debt_equity_pct', 'interest_coverage_x', 'current_ratio_pct'], '필요 원계정 미수집')
        row_notes = sorted(set(reasons.values()))
        if 'unknown' in (security, profile):
            row_notes.append('주식종류 또는 금융/비금융 분류 미확인: 일반 스크리너 대상 제외')
        companies.append({k: identity[k] for k in ('code', 'name', 'market', 'industry')} |
                         dict(security_type=security, analysis_profile=profile, metrics=metrics, history=[],
                              sources=[], data_quality=row_notes, metric_missing_reasons=reasons))
        audit[code] = {'rows': {k: len(data[k][code]) for k in data},
                       'excluded_after_completed': {k: sum(d > end for d in data[k][code]) for k in ('prices', 'flows')},
                       'financial_available_on': available, 'calculated_metrics': sum(v is not None for v in metrics.values())}
    config_hash = hashlib.sha256(json.dumps(config, ensure_ascii=False, sort_keys=True).encode()).hexdigest()
    snapshot = {'schema_version': '0.1', 'meta': {
        'price_date': end, 'flow_start': sessions[0], 'flow_end': end,
        'financial_period': period + ' 단독분기', 'financial_basis': 'CFS',
        'venue': (config['price_venue'] if config.get('price_venue') == config.get('flow_venue') and config.get('price_venue')
                  else f"가격 {config.get('price_venue') or '미확인'} / 수급 {config.get('flow_venue') or '미확인'}"),
        'universe_label': '가상 테스트 표본' if config['mode'] == 'fixture' else '인포맥스 파일 입력 표본',
        'universe_total': None, 'data_mode': config['mode'], 'as_of': as_of,
        'adapter_version': VERSION, 'config_sha256': config_hash,
        'price_venue': config.get('price_venue'), 'flow_venue': config.get('flow_venue'),
        'source_files': provenance or {}, 'warnings': warnings, 'import_audit': audit},
        'companies': sorted(companies, key=lambda r: r['code'])}
    require(all(v is None or math.isfinite(v) for r in companies for v in r['metrics'].values()), '계산값 범위 초과')
    return snapshot


def convert(config_path, paths, sheet=None):
    config = json.loads(Path(config_path).read_text(encoding='utf-8-sig'))
    grids, provenance = {}, {}
    for kind in ('info', *FIELDS):
        grids[kind], provenance[kind] = read_grid(paths[kind], sheet)
    return build_snapshot(grids, config, provenance)


def write_new(path, data):
    """Publish only a complete JSON file. Never replace an earlier result."""
    path = Path(path)
    encoded = json.dumps(data, ensure_ascii=False, indent=2, allow_nan=False).encode('utf-8')
    path.parent.mkdir(parents=True, exist_ok=True)
    require(not path.exists(), '출력 파일이 이미 있습니다. 새 이름을 사용하세요.')
    fd, temporary = tempfile.mkstemp(prefix='.infomax-', dir=path.parent)
    try:
        with os.fdopen(fd, 'wb') as out:
            out.write(encoded)
            out.flush()
            os.fsync(out.fileno())
        # Hard-link creation fails atomically when the target already exists.
        os.link(temporary, path)
    finally:
        os.unlink(temporary)


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--config', required=True, type=Path)
    for kind in ('info', *FIELDS):
        parser.add_argument('--' + kind, required=True, type=Path)
    parser.add_argument('--sheet', help='XLSX 시트명. CSV에는 적용하지 않음')
    parser.add_argument('--output', required=True, type=Path)
    args = parser.parse_args()
    try:
        result = convert(args.config, {k: getattr(args, k) for k in ('info', *FIELDS)}, args.sheet)
        write_new(args.output, result)
    except (InputError, OSError, ValueError, KeyError, BadZipFile) as exc:
        parser.exit(2, f'변환 실패: {exc}\n')
    print(f'변환 완료: {len(result["companies"])}종목 / {args.output.name}')
    print('계산 부족 사유와 파일 해시는 JSON에 기록했습니다. 공급자 자동 연결은 수행하지 않았습니다.')


if __name__ == '__main__':
    main()
